#!/usr/bin/env python3
"""Build RC Journeys Captain Property Care Agreement deliverables.

Renders the HTML source to a signature-ready PDF (headless Chromium via
Playwright) and generates an editable DOCX version of the same content.

Usage: python3 build_agreement.py
"""

import os
from pathlib import Path

HERE = Path(__file__).resolve().parent
HTML = HERE / "rc_journeys_captain_agreement.html"
PDF = HERE / "RC_Journeys_Captain_Property_Care_Agreement.pdf"
DOCX = HERE / "RC_Journeys_Captain_Property_Care_Agreement.docx"

NAVY = "1B2A4A"
GOLD = "B8912F"
RED = "8B1E1E"

FOOTER_TEMPLATE = """
<div style="width:100%; font-size:7.5pt; font-family:Arial,Helvetica,sans-serif;
            color:#5a6172; padding:0 0.65in; display:flex; justify-content:space-between;">
  <span>PT&rsquo;s Tactical Renovations LLC &nbsp;&bull;&nbsp; Veteran-Owned &nbsp;&bull;&nbsp;
        &ldquo;Your Home. Our Mission.&rdquo;</span>
  <span>Page <span class="pageNumber"></span> of <span class="totalPages"></span></span>
</div>
"""


def build_pdf():
    from playwright.sync_api import sync_playwright

    with sync_playwright() as p:
        chromium_path = os.environ.get(
            "PTTR_CHROMIUM", "/opt/pw-browsers/chromium-1194/chrome-linux/chrome"
        )
        browser = p.chromium.launch(
            executable_path=chromium_path if os.path.exists(chromium_path) else None
        )
        page = browser.new_page()
        page.goto(HTML.as_uri())
        page.pdf(
            path=str(PDF),
            format="Letter",
            print_background=True,
            display_header_footer=True,
            header_template="<div></div>",
            footer_template=FOOTER_TEMPLATE,
            margin={"top": "0.55in", "bottom": "0.8in", "left": "0", "right": "0"},
        )
        browser.close()
    print(f"PDF written: {PDF} ({PDF.stat().st_size} bytes)")


def build_docx():
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    doc = Document()

    # page setup + base style
    for section in doc.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.9)

    style = doc.styles["Normal"]
    style.font.name = "Georgia"
    style.font.size = Pt(10.5)

    def footer_text():
        footer = doc.sections[0].footer
        para = footer.paragraphs[0]
        para.text = ("PT’s Tactical Renovations LLC  •  Veteran-Owned  •  "
                     "“Your Home. Our Mission.”")
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in para.runs:
            run.font.size = Pt(7.5)
            run.font.name = "Arial"
            run.font.color.rgb = RGBColor(0x5A, 0x61, 0x72)
        # page number field on a second footer line
        p2 = footer.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p2.add_run()
        fld = run._r
        for tag, attrs, text in (
            ("w:fldChar", {"w:fldCharType": "begin"}, None),
            ("w:instrText", {"xml:space": "preserve"}, " PAGE "),
            ("w:fldChar", {"w:fldCharType": "end"}, None),
        ):
            el = fld.makeelement(qn(tag), {qn(k): v for k, v in attrs.items()})
            if text:
                el.text = text
            fld.append(el)
        run.font.size = Pt(7.5)
        run.font.name = "Arial"

    def h(text, level=1, color=NAVY):
        para = doc.add_heading(level=level)
        run = para.add_run(text)
        run.font.name = "Arial"
        run.font.color.rgb = RGBColor.from_string(color)
        run.font.size = Pt(13 if level == 1 else 11)
        return para

    def p(text, bold=False, italic=False, size=None, color=None):
        para = doc.add_paragraph()
        run = para.add_run(text)
        run.bold = bold
        run.italic = italic
        if size:
            run.font.size = Pt(size)
        if color:
            run.font.color.rgb = RGBColor.from_string(color)
        return para

    def bullets(items):
        for item in items:
            doc.add_paragraph(item, style="List Bullet")

    def checks(items, prefix="☐ "):
        for item in items:
            doc.add_paragraph(prefix + item)

    def table(rows, widths=None, header=True):
        t = doc.add_table(rows=len(rows), cols=len(rows[0]))
        t.style = "Table Grid"
        for i, row in enumerate(rows):
            for j, cell_text in enumerate(row):
                cell = t.rows[i].cells[j]
                cell.text = cell_text
                if header and i == 0:
                    for para in cell.paragraphs:
                        for run in para.runs:
                            run.bold = True
                            run.font.name = "Arial"
                            run.font.size = Pt(9)
                else:
                    for para in cell.paragraphs:
                        for run in para.runs:
                            run.font.size = Pt(9.5)
        doc.add_paragraph()
        return t

    footer_text()

    # ---------- cover ----------
    p("PT’S TACTICAL RENOVATIONS LLC  •  VETERAN-OWNED",
      bold=True, size=9, color=GOLD).alignment = WD_ALIGN_PARAGRAPH.CENTER
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("PROPERTY CARE SERVICES AGREEMENT")
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(20)
    run.font.color.rgb = RGBColor.from_string(NAVY)
    p("Captain Property Care Membership — “Your Home. Our Mission.”",
      italic=True, size=11).alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()
    table([
        ["Service Provider", "PT’s Tactical Renovations LLC — Patrick Thompson Jr., Managing Member"],
        ["Client", "R&C Journeys LLC"],
        ["Primary Property", "200 Parkview Dr., Arlington, Texas"],
        ["Membership", "Captain — $710 per month"],
        ["Effective Date", "____________________________"],
    ], header=False)

    # ---------- 1 ----------
    h("1. PARTIES, PROPERTY & PURPOSE")
    p("This Property Care Services Agreement (the “Agreement”) is entered into between "
      "PT’s Tactical Renovations LLC (“PTTR”), represented by Patrick Thompson Jr., "
      "Managing Member, and R&C Journeys LLC (the “Client”), effective as of the Effective "
      "Date stated above.")
    p("PTTR will provide membership-based handyman, property-care, and emergency-response services "
      "for the Client’s short-term rental operations at the following approved property "
      "(the “Property”):")
    p("Primary Property: 200 Parkview Dr., Arlington, Texas", bold=True)
    p("Additional properties may be added to this Agreement only by written amendment signed by both "
      "parties. Membership hours may be used only at approved properties listed in this Agreement.")

    # ---------- 2 ----------
    h("2. CAPTAIN PROPERTY CARE MEMBERSHIP")
    p("Monthly Membership Fee: $710", bold=True)
    table([
        ["Included Service", "Description"],
        ["Skilled labor hours", "12 skilled labor hours per month for work within handyman scope at the Property"],
        ["Priority scheduling", "Requests scheduled on a priority basis, generally within 24–48 hours"],
        ["Dedicated service contact", "A dedicated PTTR point of contact for all service requests"],
        ["Member emergency dispatch", "Emergency dispatch fee reduced to $50 per emergency dispatch"],
        ["Overage labor rate", "Additional labor beyond included hours billed at $75 per hour"],
        ["Hour rollover", "Unused labor hours roll over for one additional billing cycle"],
        ["Emergency stabilization", "Emergency property stabilization as defined in Section 9"],
        ["Guest-damage documentation", "Guest-damage inspection and documentation, including time-stamped "
         "photographs and written documentation suitable to support Airbnb AirCover claims"],
        ["Maintenance planning", "Routine maintenance planning for the Property"],
        ["Trade coordination", "Coordination with qualified specialty contractors when work exceeds handyman scope"],
    ])
    p("Not included in membership hours: Materials, specialty equipment, permits, disposal charges, "
      "licensed trades, subcontractors, and third-party vendor charges are NOT included in the monthly "
      "labor-hour allocation and are billed separately, unless specifically written into an approved "
      "work order.", bold=True)

    # ---------- 3 ----------
    h("3. MEMBERSHIP TERM")
    bullets([
        "This Agreement begins on the Effective Date stated in the signed contract.",
        "The membership carries an initial three-month minimum commitment.",
        "After the initial three-month term, the membership automatically converts to month-to-month.",
        "After completion of the initial term, either party may cancel with thirty (30) days’ written notice.",
        "Monthly membership fees are billed in advance.",
        "Unused labor hours may roll over for one additional billing cycle only. Expired hours have no "
        "cash value and are not refundable.",
        "Membership hours may be used only at approved R&C Journeys LLC properties listed in this Agreement.",
        "Emergency response remains subject to crew availability, site safety, weather, access, and "
        "service-area limitations.",
        "Priority response does not guarantee same-day completion.",
        "PTTR will stabilize active hazards first and will provide a separate written scope when "
        "permanent repairs exceed handyman-level work.",
    ])

    # ---------- 4 ----------
    h("4. PAYMENT STRUCTURE — SELECT ONE")
    p("☐  OPTION A — Monthly Billing", bold=True)
    bullets([
        "$710 per month",
        "Three-month initial commitment",
        "Three-month contract value: $2,130.00",
        "Charged monthly, in advance",
    ])
    p("☐  OPTION B — Three Months Paid in Advance (7% Discount)", bold=True)
    bullets([
        "Standard three-month total: $2,130.00",
        "7% prepayment discount: –$149.10",
        "Discounted three-month prepayment total: $1,980.90",
        "Effective monthly cost: $660.30",
        "Total savings: $149.10",
    ])
    h("Prepayment Conditions (Option B)", level=2)
    bullets([
        "The prepaid amount is due before the membership begins.",
        "Prepaid membership fees are nonrefundable after the service term begins, except where required by law.",
        "After the prepaid three-month period, this Agreement converts to the standard $710 "
        "month-to-month rate unless both parties approve another prepaid term in writing.",
        "The 7% discount applies only to the membership fee. It does not apply to materials, emergency "
        "dispatch charges, lawn service, overage labor, or à la carte services.",
    ])

    # ---------- 5 ----------
    h("5. ACCEPTED PAYMENT METHODS")
    bullets([
        "ACH bank transfer",
        "Business check",
        "Certified check",
        "Credit card",
        "Debit card",
        "Electronic invoice payment link",
        "Zelle, when accepted by PTTR",
        "Cash App, when accepted by PTTR",
        "Apple Pay, when accepted by PTTR",
        "Cash, only when accompanied by a written PTTR receipt",
    ])
    h("Payment Conditions", level=2)
    bullets([
        "ACH and check payments carry no processing surcharge unless a returned-payment fee applies.",
        "Credit-card and debit-card payments may include a disclosed processing fee where legally permitted.",
        "Electronic invoices may be paid through the payment processor shown on the invoice.",
        "Returned checks, rejected ACH payments, and disputed payments may result in a reasonable "
        "returned-payment or administrative fee.",
        "PTTR may suspend nonemergency services while an account is past due.",
        "Emergency response for a delinquent account is subject to payment authorization before dispatch.",
        "Payment instructions will be provided on each invoice. No bank-account numbers or sensitive "
        "financial credentials are printed in this Agreement.",
    ])
    h("Preferred Payment Method — Select One", level=2)
    checks(["ACH", "Business Check", "Credit Card", "Debit Card", "Electronic Invoice",
            "Zelle", "Cash App", "Apple Pay", "Other: __________"])

    # ---------- 6 ----------
    h("6. LAWN SERVICE ADD-ON (SEPARATE FLAT RATE)")
    p("Standard Mow and Edge Service — $65 per visit (flat rate)", bold=True)
    p("Each visit includes:")
    bullets([
        "Mowing of front, back, and side yards",
        "Edging of sidewalks, driveway, and accessible borders",
        "String-trimming around fences, structures, and obstacles",
        "Blowing clippings from sidewalks, driveway, patios, and entry areas",
    ])
    p("Lawn service is billed separately and does not deduct from the Captain Membership’s 12 "
      "included labor hours. Pricing assumes an ordinary residential lot size and routine growth. "
      "Overgrown properties, excessive debris, pet waste, storm debris, bagging, hauling, shrub work, "
      "and specialty landscaping require additional written authorization and separate pricing.")
    h("Service Frequency — Selected by Client", level=2)
    checks(["Weekly", "Every two weeks", "As requested", "Seasonal schedule", "Not selected at this time"])

    # ---------- 7 ----------
    h("7. À LA CARTE SERVICES FOR SHORT-TERM RENTAL OPERATIONS")
    p("The services below may be completed using available Captain Membership labor hours when they "
      "fall within handyman scope. Materials, specialty equipment, and services exceeding available "
      "hours are billed separately.")
    p("All à la carte services: quoted by scope, billed through available membership hours or "
      "separately authorized.", bold=True)
    menu = {
        "Guest Turnover & Interior Repairs": [
            "Drywall patching and texture blending", "Paint and trim touch-ups",
            "Door and cabinet adjustments", "Blind and curtain repairs",
            "Furniture assembly and repair", "TV and artwork mounting",
            "Shelving installation", "Caulking and cosmetic repairs", "Punch-list completion"],
        "Entry & Security": [
            "Smart-lock installation", "Lock replacement",
            "Door-latch and strike-plate adjustment", "Deadbolt replacement",
            "Keypad troubleshooting", "Temporary securing of damaged doors or windows",
            "Security-light installation"],
        "Plumbing Handyman Services": [
            "Toilet repairs", "Faucet replacement", "Minor sink repairs",
            "Garbage-disposal replacement", "Supply-line replacement",
            "Showerhead replacement", "Minor drain clearing", "Water-heater maintenance"],
        "Electrical Handyman Services": [
            "Light-fixture replacement",
            "Ceiling-fan replacement, or installation where an approved box exists",
            "Switch and receptacle replacement", "GFCI replacement",
            "Smoke-alarm installation", "Carbon-monoxide-alarm installation",
            "Exterior and security lighting"],
        "Exterior Property Services": [
            "Fence and gate repairs", "Minor siding and trim repair", "Gutter cleaning",
            "Minor gutter repair", "Pressure washing", "Exterior caulking",
            "Weatherproofing", "Exterior paint touch-ups",
            "Ground-accessible storm-debris removal"],
        "Airbnb Operations Support": [
            "Pre-arrival property inspections", "Post-checkout damage assessments",
            "Time-stamped photographs", "Written incident documentation",
            "AirCover-supporting documentation", "Emergency board-up or tarping",
            "Lockout response", "Appliance-replacement coordination",
            "Vendor-access coordination", "Property-readiness punch lists"],
        "Seasonal Services": [
            "Freeze preparation", "Storm preparation", "Storm cleanup", "Leaf cleanup",
            "Holiday-light installation and removal", "Irrigation visual inspections",
            "Seasonal exterior walkthroughs"],
    }
    for heading, items in menu.items():
        h(heading, level=2)
        bullets(items)
    p("Licensed-trade limits: Concealed leaks, sewer-line problems, major drain blockages, gas "
      "plumbing, and permit-required plumbing may require a licensed plumber. Panel work, new "
      "circuits, concealed rewiring, and permit-required electrical work require a licensed "
      "electrician. PTTR will identify these conditions and coordinate qualified licensed trades "
      "when required.", bold=True)

    # ---------- 8 ----------
    h("8. EXHIBIT A — INITIAL SERVICE AUTHORIZATION: CURRENT STORM & PRE-PHOTOGRAPHY WORK")
    h("Authorized Scope", level=2)
    bullets([
        "Remove and dispose of the fallen tree branch",
        "Assess the hanging branch",
        "Remove the hanging branch ONLY if it is safely accessible and clear of utility lines",
        "Rake and clean the affected area",
        "Patch visible cosmetic drywall or surface cracks",
        "Sand and texture-blend repaired areas",
        "Optional paint touch-up when matching paint is available",
    ])
    p("Once this Agreement becomes active, labor for this scope may be performed using the Captain "
      "Membership’s included monthly labor hours.")
    h("Billed Separately from Membership Hours", level=2)
    bullets([
        "Materials", "Debris disposal", "Specialty equipment", "Tree-service subcontractors",
        "Lift or crane service", "Utility coordination", "Work outside normal handyman scope",
    ])
    p("SAFETY EXCLUSION — UTILITY LINES: PTTR will not cut or handle branches touching, crossing, "
      "or located dangerously near utility lines. Such conditions require coordination with the "
      "utility provider or a qualified and insured tree-service contractor. PTTR will identify the "
      "condition and assist with coordination.", bold=True, color=RED)

    # ---------- 9 ----------
    h("9. EMERGENCY RESPONSE TERMS")
    table([
        ["Term", "Definition"],
        ["Priority Response", "R&C Journeys LLC is placed ahead of nonmember service requests in the "
         "scheduling queue."],
        ["Emergency Stabilization", "Immediate reasonable steps intended to stop continuing damage, "
         "remove an immediate safety hazard, or temporarily secure the Property."],
        ["Permanent Repair", "The full corrective repair, which may require a separate estimate, "
         "materials, specialty equipment, permits, or licensed trades."],
    ])
    p("Priority response does not guarantee:", bold=True)
    bullets([
        "A specific arrival time",
        "Same-day completion",
        "Immediate material availability",
        "Completion during unsafe weather",
        "Work on energized electrical systems",
        "Work involving active utility hazards",
        "Completion of licensed or permit-required work by PTTR personnel",
    ])

    # ---------- 10 ----------
    h("10. CLIENT RESPONSIBILITIES")
    p("The Client agrees to:")
    bullets([
        "Provide accurate property and access information",
        "Identify active guests, pets, alarms, and access restrictions before each visit",
        "Maintain safe access to the work area",
        "Disclose known hazards and utility conditions",
        "Provide written authorization for additional work",
        "Pay invoices according to the selected payment schedule",
        "Inform PTTR of emergency conditions promptly",
        "Maintain appropriate property and liability insurance",
    ])

    # ---------- 11 ----------
    doc.add_page_break()
    h("11. ACCEPTANCE & SIGNATURES")
    p("By signing below, each party accepts the terms of this Agreement, including the selected "
      "payment structure, payment method, and lawn-service frequency indicated.")
    h("Client — R&C Journeys LLC", level=2)
    p("Authorized Representative Name: ________________________________________")
    p("Title: ________________________________________")
    p("Signature: ________________________________________        Date: ______________")
    h("Selected Payment Option", level=2)
    checks(["Option A — $710 monthly",
            "Option B — $1,980.90 prepaid for the initial three months"])
    p("Selected Payment Method: ________________________________________")
    h("Lawn Service Selection", level=2)
    checks(["Weekly", "Every two weeks", "As requested", "Seasonal", "Not selected"])
    h("Service Provider — PT’s Tactical Renovations LLC", level=2)
    p("Patrick Thompson Jr., Managing Member")
    p("Signature: ________________________________________        Date: ______________")
    h("Electronic Signature Clause", level=2)
    p("The parties agree that electronic signatures, digital acceptance and signed PDF copies are "
      "enforceable to the same extent as original handwritten signatures. Each signer confirms that "
      "they are authorized to enter into this agreement on behalf of the organization identified above.")

    doc.save(str(DOCX))
    print(f"DOCX written: {DOCX} ({DOCX.stat().st_size} bytes)")


if __name__ == "__main__":
    build_pdf()
    build_docx()
